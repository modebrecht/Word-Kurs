from __future__ import annotations

"""Optionale Transferaufträge V3: keine A14, keine Notenänderung."""
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from course_common import base_doc, block, add_text, add_step, add_tip, add_check, add_finish, finalise, NAVY, TEAL, TEAL_DARK, PALE_TEAL, WHITE, set_run


def _raw_style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    doc.core_properties.subject = "Word-Kurs · Sek 8 · Praxisdokument"
    doc.core_properties.keywords = "Word-Kurs; Sekundarstufe I; Sek 8; Informatikbildung"
    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(17)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(NAVY)
    heading.paragraph_format.space_after = Pt(8)


def _add_link(paragraph, title: str, target: str) -> None:
    """Create a real external Word hyperlink, not blue-looking plain text."""
    rid = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL_DARK)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    txt = OxmlElement("w:t")
    txt.text = title
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)


def _instruction_brief(out: Path) -> None:
    doc = base_doc(
        "PLUS 1", "Schnupperlehre", "Ein Brief, der ankommt",
        "Gestalte aus vorgegebenem Rohtext eine professionelle Anfrage.",
        "einen formellen Brief strukturieren, sauber speichern und nur die Briefseite als PDF exportieren.",
    )
    t = block(doc, "01", "VORBEREITEN")
    cell = t.cell(0, 1)
    add_text(cell, "Öffne START_Schnupperlehrbrief.docx. Speichere sofort eine eigene Kopie als Schnupperbrief_Nachname.docx im Ordner IB.", 9.5)
    t = block(doc, "02", "GESTALTEN")
    cell = t.cell(0, 1)
    add_text(cell, "Alle Texte sind vorhanden. Die Firma und die Person sind fiktiv. Sende diesen Übungsbrief nicht ab.", 9.3)
    add_text(cell, "BRIEFSCHEMA · So liegen die Teile auf der Briefseite:", 9.2, bold=True, color=NAVY, after=.2)
    schema = cell.add_table(rows=4, cols=2)
    for row_index, (left_text, right_text) in enumerate([
        ("Absender: Name und Adresse", ""),
        ("Empfänger: Firma und Adresse", "Ort, Datum"),
        ("Betreff (fett) · Anrede", ""),
        ("Zwei Textabsätze · Gruss · Name", ""),
    ]):
        for col_index, label in enumerate((left_text, right_text)):
            p = schema.cell(row_index, col_index).paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_index == 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(label)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.bold = row_index in (0, 2)

    add_step(cell, "A", [("A4 Hochformat, Ränder Normal, Arial 11 pt.", True, False, NAVY)])
    add_step(cell, "B", [("Absender oben, Empfänger darunter, Datum rechts, dann Betreff.", True, False, NAVY)])
    add_step(cell, "C", [("Betreff fett; Anrede, zwei kurze Textabsätze und Grussformel gut lesbar.", True, False, NAVY)])
    add_step(cell, "D", [("Keine Abstände mit vielen Leerzeichen erzeugen. Nutze Absätze und Absatzabstände.", True, False, NAVY)])
    t = block(doc, "03", "EXPORTIEREN")
    cell = t.cell(0, 1)
    add_text(cell, "1 · Nach dem Gestalten zuerst Strg + S: fertige DOCX-Datei speichern. Druckvorschau: genau eine Briefseite kontrollieren.", 9.25)
    add_text(cell, "2 · PDF aus der BRIEFDATEI exportieren: Datei → Speichern unter → Dateityp PDF (oder Datei → Exportieren). Name: Schnupperbrief_Nachname.pdf.", 9.25)
    add_text(cell, "3 · PDF öffnen und mit der erneut geöffneten DOCX vergleichen: gleicher Text, gleiche Anordnung, genau eine Seite. Beide Dateien müssen im Ordner IB liegen.", 9.25)
    add_tip(doc, "DOCX = später bearbeitbar. PDF = zum Weitergeben, ohne dass sich das Layout leicht verschiebt.", "MERKE")
    add_check(doc, "Stimmt der Briefaufbau? Wurde die DOCX nach der Bearbeitung gespeichert? Sind die erneut geöffneten DOCX und PDF inhaltlich und optisch gleich, jeweils genau 1 Seite?")
    add_finish(doc, "Gib die DOCX- und die PDF-Datei im Ordner IB ab. Den Brief nicht an ein Unternehmen senden.")
    finalise(doc, out, "Zusatz: Schnupperlehrbrief")


def _starter_brief(out: Path) -> None:
    doc = Document()
    _raw_style(doc)
    doc.core_properties.title = "Rohtext: Schnupperlehrbrief"
    lines = [
        "Sara Muster", "Beispielweg 4", "4500 Solothurn",
        "Pixelwerkstatt AG", "Berufsbildung", "Werkweg 12", "4500 Solothurn",
        "Solothurn, 14. Oktober 2026",
        "Anfrage für eine Schnupperlehre als Informatikerin EFZ",
        "Sehr geehrte Damen und Herren",
        "Ich interessiere mich für den Beruf Informatikerin EFZ. Darum möchte ich gerne einen Einblick in den Berufsalltag erhalten.",
        "Über eine Möglichkeit zum Schnuppern würde ich mich freuen. Ich bin motiviert und lerne gern Neues.",
        "Freundliche Grüsse", "Sara Muster",
    ]
    for item in lines:
        doc.add_paragraph(item)
    doc.save(out)


def _model_brief(out: Path) -> None:
    doc = Document()
    _raw_style(doc)
    doc.core_properties.title = "Musterlösung: Schnupperlehrbrief"
    for s in ["Sara Muster", "Beispielweg 4", "4500 Solothurn"]:
        p = doc.add_paragraph(s)
        p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph("")
    p.paragraph_format.space_after = Pt(4)
    for s in ["Pixelwerkstatt AG", "Berufsbildung", "Werkweg 12", "4500 Solothurn"]:
        p = doc.add_paragraph(s)
        p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph("Solothurn, 14. Oktober 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(15)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Anfrage für eine Schnupperlehre als Informatikerin EFZ")
    r.bold = True
    doc.add_paragraph("Sehr geehrte Damen und Herren")
    doc.add_paragraph("Ich interessiere mich für den Beruf Informatikerin EFZ. Darum möchte ich gerne einen Einblick in den Berufsalltag erhalten.")
    doc.add_paragraph("Über eine Möglichkeit zum Schnuppern würde ich mich freuen. Ich bin motiviert und lerne gern Neues.")
    doc.add_paragraph("Freundliche Grüsse")
    doc.add_paragraph("Sara Muster")
    doc.save(out)


def _instruction_links(out: Path) -> None:
    doc = base_doc(
        "PLUS 2", "Links & Quellen", "Links mit Sinn",
        "Verbinde verständliche Linktexte mit richtigen Webseiten und nenne eine Quelle.",
        "einen echten Word-Hyperlink setzen und eine einfache Quellenangabe schreiben.",
    )
    t = block(doc, "01", "ZWEI LINKS")
    cell = t.cell(0, 1)
    add_text(cell, "Öffne START_Links_und_Quellen.docx und speichere eine Kopie im Ordner IB.", 9.5)
    add_step(cell, "A", [("Markiere «Lehre EFZ (Berufsberatung Schweiz)» und drücke Strg + K.", True, False, NAVY), ("  Adresse: https://www.berufsberatung.ch/de/bildungssystem/d/lehre-efz", False, False, None)])
    add_step(cell, "B", [("Markiere «SBB Fahrplan» und drücke Strg + K.", True, False, NAVY), ("  Adresse: https://www.sbb.ch/", False, False, None)])
    t = block(doc, "02", "FAKT + QUELLE")
    cell = t.cell(0, 1)
    add_text(cell, "Öffne den ersten Link (in Word meist Strg + Klick). Lies auf der Seite «Lehre EFZ» den Abschnitt «Dauer». Ergänze im Startdokument in eigenen Worten, wie viele Jahre eine EFZ-Lehre dauert.", 9.35)
    add_text(cell, "Trage bei «Quelle:» Anbieter, Seitentitel «Lehre EFZ», genaue Webadresse und heutiges Zugriffsdatum ein. So belegt die Quelle eine konkrete Aussage.", 9.35)
    add_tip(doc, "Ein guter Linktext verrät das Ziel. Eine Quelle muss zur verwendeten Information passen, nicht irgendeine Seite nennen. Einem Link in Word folgst du meist mit Strg + Klick.", "MERKE")
    add_check(doc, "Sind zwei echte Hyperlinks vorhanden? Stimmt die Dauer der EFZ-Lehre mit der verlinkten Seite überein? Belegt die Quellenzeile diese Information mit heutigem Zugriffsdatum?")
    add_finish(doc, "Speichere die bearbeitete DOCX im Ordner IB. Dieser Zusatzauftrag ist freiwillig und nicht Teil des Word-Tests.")
    finalise(doc, out, "Zusatz: Links und Quellen")


def _starter_links(out: Path) -> None:
    doc = Document()
    _raw_style(doc)
    doc.core_properties.title = "Rohtext: Links und Quellen"
    doc.add_heading("Hilfreiche Webseiten für Schule und Beruf", 1)
    doc.add_paragraph("Für Infos zur EFZ-Lehre: Lehre EFZ (Berufsberatung Schweiz)")
    doc.add_paragraph("Für Zugverbindungen: SBB Fahrplan")
    doc.add_paragraph("Das habe ich herausgefunden: ")
    doc.add_paragraph("Quelle: ")
    doc.save(out)


def _model_links(out: Path) -> None:
    doc = Document()
    _raw_style(doc)
    doc.core_properties.title = "Musterlösung: Links und Quellen"
    doc.add_heading("Hilfreiche Webseiten für Schule und Beruf", 1)
    p = doc.add_paragraph("Für Infos zur EFZ-Lehre: ")
    _add_link(p, "Lehre EFZ (Berufsberatung Schweiz)", "https://www.berufsberatung.ch/de/bildungssystem/d/lehre-efz")
    p = doc.add_paragraph("Für Zugverbindungen: ")
    _add_link(p, "SBB Fahrplan", "https://www.sbb.ch/")
    doc.add_paragraph("Das habe ich herausgefunden: Eine EFZ-Lehre dauert 3 oder 4 Jahre.")
    doc.add_paragraph("Quelle: Berufsberatung Schweiz, «Lehre EFZ», www.berufsberatung.ch/de/bildungssystem/d/lehre-efz, Zugriff am [Datum der Bearbeitung].")
    doc.save(out)


def _teacher(out: Path) -> None:
    doc = base_doc(
        "LP V3", "Lehrpersonenmaterial", "Unterrichtshinweise",
        "Die Ergänzungen zu A1, A7 und den freiwilligen PLUS-Aufträgen.",
        "Die Pflichtfolge A1–A13 und alle drei bestehenden Noten bleiben unverändert. "
        "PLUS 1 und PLUS 2 sind freiwillig und geben keine Fleisspunkte.",
        header_kind="LEHRPERSON", goal_caption="ÜBERBLICK",
    )
    doc.core_properties.title = "Lehrpersonenhilfen: Word-Kurs V3"
    doc.core_properties.subject = "Word-Kurs · Sek 8 · Lehrpersonenmaterial"

    sections = [
        (
            "A1 · Sicher speichern (5–8 Minuten)",
            "VOR der Formatierung: Datei → Speichern unter → persönlicher Ordner IB; Namen vergeben. "
            "NACH der Bearbeitung: Strg + S, schliessen und aus IB erneut öffnen. "
            "Schwächere Lernende bekommen den Speicherort gezeigt und führen den Ablauf einmal vor; "
            "fortgeschrittene Lernende erledigen ihn selbst. Formative Evidenz: die fertige Datei lässt sich wiederfinden. "
            "Kein Abzug für zunächst falsches Speichern, wenn anschliessend korrigiert."
        ),
        (
            "A7 · Bild: informativ oder dekorativ? (freiwillig, 5 Minuten)",
            "Erst das Bildlayout fertigstellen. Beim gezeichneten Schulhaus kann die Illustration im vorhandenen Text "
            "lediglich schmücken: dann «Als dekorativ kennzeichnen». "
            "Ein informatives Bild vermittelt relevante neue Inhalte: nur diese Information knapp als Alternativtext "
            "beschreiben. Die Unterscheidung ist Lernziel; das blosse Abschreiben einer Bildbeschreibung nicht. "
            "Word-Funktion vor Ort prüfen; kein Prüfungsstoff."
        ),
        (
            "PLUS 1 · Schnupperlehrbrief und PDF (Richtwert 45 Minuten)",
            "5 Minuten Briefschema erklären (Absender/Empfänger links, Datum rechts, Betreff/Anrede/Text/Gruss). "
            "25 Minuten vorhandenen Rohtext ordnen, 10 Minuten Strg + S / PDF-Export / BEIDE Dateien öffnen, "
            "5 Minuten Reserve. Unterlegen: START_Schnupperlehrbrief.docx und Aufgabenblatt. "
            "Schwächere Lernende orientieren sich an den Schema-Zeilen; andere arbeiten ohne weiteren Klickplan. "
            "Formativ: korrekte Bestandteile, eine Seite, identisches DOCX/PDF, Ordner IB. "
            "Die Musterlösung bleibt bei der Lehrperson; Übungsbrief nicht real versenden."
        ),
        (
            "PLUS 2 · Hyperlinks, Aussage und Quellenbeleg (25–30 Minuten)",
            "Strg + K zum Einfügen eines Links; in Word einem Link normalerweise mit Strg + Klick folgen. "
            "Erster Link führt direkt zu «Lehre EFZ» bei berufsberatung.ch; dort unter «Dauer» die konkrete "
            "Information finden: EFZ-Lehre dauert drei oder vier Jahre. Diesen Fakt in einem eigenen Satz "
            "notieren und mit Anbieter, Titel, Seitenadresse und Zugriffsdatum belegen. Zweiter Link: SBB Fahrplan. "
            "Schwächere Lernende erhalten die Seite geöffnet; falls Schulfilter sperren, einen Ausdruck der "
            "Originalquelle zur Verfügung stellen. Ein festes Musterdatum darf nicht abgeschrieben werden: "
            "die Musterlösung enthält bewusst [Datum der Bearbeitung]."
        ),
        (
            "A12/A13 und benotete Leistungsnachweise",
            "Beide Pflichtblätter sowie alle drei bestehenden Noten unverändert lassen. "
            "PDF-Export, Alternativtext oder Linkwissen nicht überraschend als zusätzliche "
            "Prüfungskompetenz verlangen. Zuerst Pflichtauftrag, dann freiwillige PLUS-Aufgaben."
        ),
    ]
    for heading, text in sections:
        if " · " in heading:
            code, subtitle = heading.split(" · ", 1)
        else:
            code, subtitle = "A12/A13", "Leistungsnachweise"
        section = block(doc, code, subtitle, label_size=10.0)
        paragraph = section.cell(0, 1).paragraphs[0]
        paragraph.paragraph_format.line_spacing = 1.08
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(text)
        set_run(run, size=9.45)
    check = block(
        doc, "CHECK", "Beobachtung", fill_left=PALE_TEAL,
        fill_right=PALE_TEAL, label_color=TEAL_DARK, label_size=11.0,
    )
    cell = check.cell(0, 1)
    for item in [
        "A1: vor Arbeitsbeginn speichern; nach Änderungen Strg + S; Datei im korrekten Ordner erneut öffnen.",
        "A7: dekorativ vs. informativ begründet entscheiden, ohne die Bildlayout-Hauptaufgabe zu verändern.",
        "PLUS 1: Briefschema, gespeicherte End-DOCX, geöffnete PDF und visuelle Übereinstimmung.",
        "PLUS 2: zwei echte Hyperlinks, korrekte Dauer-Aussage, passende Quelle, aktuelles Datum.",
        "Bei Schwierigkeiten: Bedienung sichtbar vormachen; keine versteckten Zusatzanforderungen aufbauen.",
        "Extern offen: Microsoft Word 365 auf Schul-PCs und echter Schülerdurchlauf samt Zeitmessung.",
    ]:
        add_text(cell, "• " + item, size=9.2, after=.35)
    doc.save(out)


def build_all(root: Path) -> None:
    output = Path(root) / "arbeitsblaetter"
    teacher = output / "lehrperson"
    teacher.mkdir(parents=True, exist_ok=True)
    _instruction_brief(output / "ZUSATZ_Schnupperlehrbrief.docx")
    _starter_brief(output / "START_Schnupperlehrbrief.docx")
    _model_brief(teacher / "MUSTER_Schnupperlehrbrief.docx")
    _instruction_links(output / "ZUSATZ_Links_und_Quellen.docx")
    _starter_links(output / "START_Links_und_Quellen.docx")
    _model_links(teacher / "MUSTER_Links_und_Quellen.docx")
    _teacher(teacher / "LEHRPERSONENHILFEN_V3.docx")
